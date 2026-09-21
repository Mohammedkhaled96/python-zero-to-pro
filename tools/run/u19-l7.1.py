from openpyxl import Workbook, load_workbook
wb = Workbook()
ws = wb.active
ws.title = "الرواتب"
ws.append(["الاسم", "الراتب", "بعد الزيادة"])
for row, (name, salary) in enumerate([("سارة", 12000), ("أحمد", 9500), ("منى", 11000)], start=2):
    ws.append([name, salary, f"=B{row}*1.1"])
wb.save("salaries.xlsx")
sheet = load_workbook("salaries.xlsx")["الرواتب"]
for name, salary, formula in sheet.iter_rows(min_row=2, values_only=True):
    print(name, salary, formula)
